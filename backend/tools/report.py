"""自动生成企业经营数据分析 Word 报告。"""
from __future__ import annotations
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from langchain_core.tools import tool
from langchain_openai import ChatOpenAI
from .anomaly import detect_zscore_anomalies
from .business_kpi import calculate_business_kpis
from .correlation import analyze_correlation
from .descriptive import descriptive_statistics
from .quality import check_data_quality
from .trend import analyze_time_trend
from .visualization import generate_business_charts
from .common import read_dataframe

#定义报告生成格式
BLUE = RGBColor(46, 116, 181); DARK = RGBColor(31, 77, 120)


def _font(run, size=12, bold=False, color=None, font_name="宋体"):
    run.font.name = "SimSun"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "SimSun")
    fonts.set(qn("w:hAnsi"), "SimSun")
    fonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = color


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)

#调用模型以及模型要求
def _create_report_llm() -> ChatOpenAI:
    """创建专门用于报告解读的阿里云 Qwen 客户端。"""
    api_key = os.getenv("DASHSCOPE_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("未设置 DASHSCOPE_API_KEY，无法生成 AI 解读")
    model_name = os.getenv("DASHSCOPE_MODEL", "qwen2.5-14b-instruct")
    kwargs: dict[str, Any] = {}
    if model_name.lower().startswith("qwen3"):
        kwargs["extra_body"] = {"enable_thinking": False}
    return ChatOpenAI(
        model=model_name,
        base_url=os.getenv("DASHSCOPE_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1"),
        api_key=api_key,
        temperature=0.2,
        timeout=180,
        max_retries=3,
        max_tokens=380,
        **kwargs,
    )

#prompt设计及处理模型输出
def _ask_qwen(model: ChatOpenAI, subject: str, evidence: dict[str, Any]) -> str:
    """只根据压缩后的统计证据生成约 200 字中文解读。"""
    prompt = f"""你是企业经营数据分析师。请解读“{subject}”。
只能使用下方统计证据，不要重新计算，不要补充证据中没有的原因。
写成一段180至230个中文字符的简洁分析，不使用Markdown，不列项目符号。
先说明数据表现，再说明业务含义或风险；相关性分析必须注明相关不代表因果。
统计证据：{json.dumps(evidence, ensure_ascii=False, default=str)}"""
    content = model.invoke(prompt).content
    if isinstance(content, list):
        content = "".join(str(item.get("text", item)) if isinstance(item, dict) else str(item) for item in content)
    text = re.sub(r"\s+", "", str(content)).strip()
    return text[:260]


def _add_ai_interpretation(doc: Document, text: str) -> None:
    """使用段落底纹插入与图表相邻的 AI 解读，不用表格包装正文。"""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.2
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F6FA")
    p_pr.append(shading)
    _font(paragraph.add_run("AI 解读："), 12, True, DARK)
    _font(paragraph.add_run(text), 12)


def _table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"; table.autofit = False
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = str(text); _shade(cell, "E8EEF5"); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs: _font(run, 10, True, DARK)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = "-" if value is None else str(value); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[i].paragraphs[0].runs: _font(run, 9.5)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths): cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _setup(doc):
    section = doc.sections[0]; section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    normal = doc.styles["Normal"]; normal.font.name = "SimSun"; normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.1
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "SimSun"
        style.font.size = Pt(12)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    for name, size, color, before, after in (("Heading 1",16,BLUE,16,8),("Heading 2",13,BLUE,12,6),("Heading 3",12,DARK,8,4)):
        style = doc.styles[name]; style.font.name = "Microsoft YaHei"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]; header.text = "企业经营数据分析报告"; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs: _font(run, 9, False, RGBColor(100,100,100))

#引用tool完成报告书写，分模块解释
@tool
def generate_word_report(file_path: str, output_path: str, company_name: str = "示例企业",
                         date_column: str | None = None, sales_column: str | None = None,
                         profit_column: str | None = None, quantity_column: str | None = None,
                         order_id_column: str | None = None, group_column: str | None = None,
                         advertising_column: str | None = None,
                         use_llm_interpretation: bool = True) -> dict[str, Any]:
    """汇总数据质量、统计、KPI、趋势、相关性、异常和图表，生成 DOCX 报告。"""
    target = Path(output_path).resolve(); target.parent.mkdir(parents=True, exist_ok=True)
    charts_dir = target.parent / f"{target.stem}_charts"
    quality = check_data_quality.invoke({"file_path": file_path})
    numeric = [c for c in (sales_column, profit_column, quantity_column, advertising_column) if c]
    desc = descriptive_statistics.invoke({"file_path": file_path, "columns": numeric or None})
    kpis = calculate_business_kpis.invoke({"file_path": file_path, "sales_column": sales_column,
        "profit_column": profit_column, "quantity_column": quantity_column, "order_id_column": order_id_column}) if sales_column else {}
    trend = analyze_time_trend.invoke({"file_path": file_path, "date_column": date_column,
        "value_column": sales_column, "frequency": "month"}) if date_column and sales_column else None
    anomaly = detect_zscore_anomalies.invoke({"file_path": file_path, "value_column": sales_column,
        "date_column": date_column, "threshold": 3.0, "frequency": "month"}) if date_column and sales_column else None
    correlation = analyze_correlation.invoke({"file_path": file_path, "x_column": advertising_column,
        "y_column": sales_column}) if advertising_column and sales_column else None
    chart_args = {"file_path": file_path, "output_dir": str(charts_dir), "date_column": date_column,
                  "sales_column": sales_column, "group_column": group_column,
                  "x_column": advertising_column, "y_column": sales_column if advertising_column else None}
    charts = generate_business_charts.invoke(chart_args)["charts"] if any((date_column and sales_column, group_column and sales_column, advertising_column and sales_column)) else {}

    group_summary = None
    if group_column and sales_column:
        frame = read_dataframe(file_path)
        # CSV 中的数值列经常被 Pandas 识别为 object/string。直接 groupby.sum()
        # 会把 "71.37"、"23083.43" 拼成超长字符串，之后 float() 必然失败。
        numeric_sales = pd.to_numeric(frame[sales_column], errors="coerce")
        valid_group_rows = pd.DataFrame({
            "_group": frame[group_column],
            "_sales": numeric_sales,
        }).dropna(subset=["_group", "_sales"])
        if valid_group_rows.empty:
            raise ValueError(f"字段“{sales_column}”没有可用于分组分析的有效数值。")
        grouped = (
            valid_group_rows.groupby("_group", dropna=True)["_sales"]
            .sum()
            .sort_values(ascending=False)
        )
        group_summary = {
            "group_column": group_column,
            "sales_column": sales_column,
            "top_groups": [{"group": str(k), "sales": round(float(v), 2)} for k, v in grouped.head(10).items()],
            "group_count": int(len(grouped)),
        }

    interpretations: dict[str, str] = {}
    interpretation_errors: list[str] = []
    if use_llm_interpretation:
        try:
            report_llm = _create_report_llm()
            trend_evidence = None if not trend else {
                "direction": trend["direction"],
                "overall_growth_percent": trend["overall_growth_percent"],
                "peak_period": trend["peak_period"],
                "peak_value": trend["peak_value"],
                "lowest_period": trend["lowest_period"],
                "lowest_value": trend["lowest_value"],
                "first_period": trend["periods"][0] if trend["periods"] else None,
                "last_period": trend["periods"][-1] if trend["periods"] else None,
            }
            evidence_map = {
                "statistics": {"data_quality": {"row_count": quality["row_count"], "duplicate_rows": quality["duplicate_rows"],
                    "problem_columns": quality["problem_columns"]}, "descriptive_statistics": desc, "business_kpis": kpis},
                "trend": trend_evidence,
                "group": group_summary,
                "scatter": correlation,
                "boxplot": anomaly,
                "management": {"business_kpis": kpis, "trend": trend_evidence, "correlation": correlation,
                    "anomalies": anomaly, "group_summary": group_summary},
            }
            subjects = {"statistics": "描述统计与业务KPI", "trend": "销售趋势图", "group": "分组销售柱状图",
                        "scatter": "相关性散点图", "boxplot": "月度箱线图与异常点", "management": "管理建议"}
            for key, evidence in evidence_map.items():
                if evidence:
                    try:
                        interpretations[key] = _ask_qwen(report_llm, subjects[key], evidence)
                    except Exception as exc:
                        interpretation_errors.append(f"{key}: {type(exc).__name__}: {exc}")
        except Exception as exc:
            interpretation_errors.append(f"model: {type(exc).__name__}: {exc}")

    doc = Document(); _setup(doc)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(110)
    _font(p.add_run("企业经营数据分析报告"), 28, True, DARK, "微软雅黑")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _font(p.add_run(company_name), 16, False, BLUE, "微软雅黑")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(30)
    _font(p.add_run(f"数据来源：{Path(file_path).name}\n生成时间：{datetime.now():%Y-%m-%d %H:%M}"), 10, False, RGBColor(90,90,90))
    doc.add_page_break()

    doc.add_heading("1. 执行摘要", level=1)
    summary = [f"数据共 {quality['row_count']:,} 行；重复记录 {quality['duplicate_rows']} 行。"]
    if kpis: summary.append(f"总销售额 {kpis['total_sales']:,.2f}；平均订单金额 {kpis['average_order_value']:,.2f}。")
    if "profit_margin_percent" in kpis: summary.append(f"总利润 {kpis['total_profit']:,.2f}；利润率 {kpis['profit_margin_percent']:.2f}%。")
    if trend: summary.append(f"销售趋势总体{trend['direction']}，首末期变化 {trend['overall_growth_percent']}%。")
    if anomaly: summary.append(f"Z-score 检测发现 {anomaly['anomaly_count']} 个异常期间。")
    for text in summary: doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("2. 数据概览与质量", level=1)
    _table(doc, ["指标", "结果"], [["数据行数", quality["row_count"]], ["重复行", quality["duplicate_rows"]],
        ["存在质量问题", "是" if quality["has_quality_issues"] else "否"], ["缺失字段数", len(quality["problem_columns"])]], [2.2,4.3])
    if quality["problem_columns"]:
        _table(doc, ["字段", "缺失数", "缺失率"], [[c,v["missing_count"],f"{v['missing_rate_percent']}%"] for c,v in quality["problem_columns"].items()], [2.8,1.5,2.2])

    doc.add_heading("3. 描述统计与业务 KPI", level=1)
    if desc["statistics"]:
        _table(doc, ["字段","均值","中位数","标准差","最小值","最大值"],
               [[c,v["mean"],v["median"],v["std"],v["min"],v["max"]] for c,v in desc["statistics"].items()], [1.4,1,1,1.1,1,1])
    if kpis: _table(doc, ["KPI", "数值"], [[k.replace("_"," "), v] for k,v in kpis.items()], [3.1,3.4])
    if "statistics" in interpretations:
        _add_ai_interpretation(doc, interpretations["statistics"])

    doc.add_heading("4. 趋势、相关性与异常", level=1)
    if trend:
        doc.add_paragraph(f"峰值期间：{trend['peak_period']}（{trend['peak_value']:,.2f}）；最低期间：{trend['lowest_period']}（{trend['lowest_value']:,.2f}）。")
    if correlation:
        doc.add_paragraph(f"{advertising_column} 与 {sales_column} 的 Pearson r={correlation['pearson_r']}，表现为{correlation['strength']}{correlation['direction']}，p={correlation['p_value']}。相关性不代表因果关系。")
    if anomaly:
        _table(doc, ["异常期间", "数值", "Z-score", "方向"], [[a["label"],a["value"],a["z_score"],a["direction"]] for a in anomaly["anomalies"]] or [["未发现",None,None,None]], [2.2,1.5,1.4,1.4])
    for title, path in charts.items():
        doc.add_heading({"trend":"销售趋势图","group":"分组销售图","scatter":"相关性散点图","boxplot":"月度箱线图"}.get(title,title), level=2)
        doc.add_picture(path, width=Inches(6.1)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if title in interpretations:
            _add_ai_interpretation(doc, interpretations[title])

    doc.add_heading("5. 管理建议", level=1)
    recommendations = ["优先核查异常期间的区域、品类和订单变化，避免仅凭总体指标判断原因。",
                       "对低利润或亏损记录按区域与品类进一步拆分，识别产品结构和折扣影响。",
                       "将月度 KPI、异常阈值和数据质量指标纳入固定经营复盘。"]
    if correlation and correlation["strength"] in ("中等","强"):
        recommendations.append("广告投入与销售存在统计相关，可设计分组实验进一步验证因果效果。")
    for item in recommendations: doc.add_paragraph(item, style="List Number")
    if "management" in interpretations:
        _add_ai_interpretation(doc, interpretations["management"])
    if interpretation_errors:
        note = doc.add_paragraph("部分 AI 解读生成失败，报告保留了确定性统计结果。")
        for run in note.runs: _font(run, 9, False, RGBColor(155, 28, 28))
    doc.save(target)
    web_summary = interpretations.get("management") or " ".join(summary)
    return {"report_path": str(target), "chart_paths": charts, "sections": 5,
            "quality_issues": quality["has_quality_issues"], "anomaly_count": anomaly["anomaly_count"] if anomaly else None,
            "llm_interpretation_count": len(interpretations), "llm_interpretation_errors": interpretation_errors,
            "summary": web_summary}
